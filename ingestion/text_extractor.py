from pypdf import PdfReader
import docx
import io
import logging
import warnings

try:
    from pptx import Presentation
except ImportError:
    Presentation = None

try:
    from PIL import Image
    from PIL import ImageFile
except ImportError:
    Image = None
    ImageFile = None

try:
    import pytesseract
except ImportError:
    pytesseract = None


# Reduce noisy third-party parser logs during bulk indexing.
logging.getLogger("pypdf").setLevel(logging.ERROR)
warnings.filterwarnings(
    "ignore",
    message=r".*quantization tables are too coarse for baseline JPEG.*",
)

if ImageFile is not None:
    # Avoid hard failures on partially damaged JPEG/PNG files.
    ImageFile.LOAD_TRUNCATED_IMAGES = True


def extract_text(file_data: dict, use_ocr: bool = True):
    """
    Extract text from different file types.
    """

    file_path = file_data["file_path"]
    file_extension = file_data["file_extension"]

    if file_extension == ".pdf":
        return extract_pdf_text(file_path, use_ocr=use_ocr)

    elif file_extension == ".txt":
        return extract_txt_text(file_path)

    elif file_extension == ".docx":
        return extract_docx_text(file_path)

    elif file_extension == ".pptx":
        return extract_pptx_text(file_path, use_ocr=use_ocr)

    elif file_extension in {".png", ".jpg", ".jpeg"}:
        return extract_image_text(file_path, use_ocr=use_ocr)

    else:
        raise ValueError(f"Unsupported file type: {file_extension}")


def extract_pdf_text(file_path, use_ocr: bool = True):
    if not _looks_like_pdf(file_path):
        raise ValueError("Invalid PDF header")

    try:
        reader = PdfReader(file_path, strict=False)
    except Exception as exc:
        raise ValueError(f"Unreadable PDF: {exc}") from exc

    text_parts = []

    for page in reader.pages:
        page_text = (page.extract_text() or "").strip()

        if page_text:
            text_parts.append(page_text)
            continue

        # OCR fallback for scanned/image-heavy PDF pages.
        if not use_ocr:
            continue

        ocr_text = _extract_text_from_pdf_page_images(page)
        if ocr_text:
            text_parts.append(ocr_text)

    return "\n".join(text_parts)


def _extract_text_from_pdf_page_images(page) -> str:
    """
    Attempt OCR on images embedded inside a PDF page.
    """

    if Image is None or pytesseract is None:
        return ""

    page_images = getattr(page, "images", None)
    if not page_images:
        return ""

    image_text_parts = []

    for pdf_image in page_images:
        image_bytes = getattr(pdf_image, "data", None)
        if not image_bytes:
            continue

        image_text = _extract_text_from_image_bytes(image_bytes).strip()
        if image_text:
            image_text_parts.append(image_text)

    return "\n".join(image_text_parts)


def extract_txt_text(file_path):
    with open(file_path, "r", encoding="utf-8") as f:
        return f.read()


def extract_docx_text(file_path):
    doc = docx.Document(file_path)

    text = []
    for paragraph in doc.paragraphs:
        text.append(paragraph.text)

    return "\n".join(text)


def extract_pptx_text(file_path, use_ocr: bool = True):
    if Presentation is None:
        raise ImportError(
            "python-pptx is required for .pptx files. Install it with: pip install python-pptx"
        )

    presentation = Presentation(file_path)
    text_parts = []

    for slide in presentation.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text") and shape.text and shape.text.strip():
                text_parts.append(shape.text.strip())

            if use_ocr and hasattr(shape, "image"):
                ocr_text = _extract_text_from_image_bytes(shape.image.blob)
                if ocr_text:
                    text_parts.append(ocr_text)

    return "\n".join(text_parts)


def extract_image_text(file_path, use_ocr: bool = True):
    if not use_ocr:
        return ""

    if Image is None or pytesseract is None:
        raise ImportError(
            "Pillow and pytesseract are required for image OCR. Install with: pip install pillow pytesseract"
        )

    try:
        with Image.open(file_path) as image:
            return pytesseract.image_to_string(image)
    except Exception as exc:
        raise ValueError(f"Unreadable image: {exc}") from exc


def _extract_text_from_image_bytes(image_bytes: bytes) -> str:
    if Image is None or pytesseract is None:
        return ""

    try:
        with Image.open(io.BytesIO(image_bytes)) as image:
            return pytesseract.image_to_string(image)
    except Exception:
        # Skip bad embedded images instead of failing full file indexing.
        return ""


def _looks_like_pdf(file_path: str) -> bool:
    try:
        with open(file_path, "rb") as file_obj:
            header = file_obj.read(5)
        return header == b"%PDF-"
    except Exception:
        return False